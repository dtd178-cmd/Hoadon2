"""
Module phân tích file Excel mẫu và ghi dữ liệu hóa đơn vào đúng sheet, đúng cột
mà vẫn bảo toàn cấu trúc, công thức, định dạng và các sheet khác của file mẫu.
"""
from copy import copy
from pathlib import Path
from typing import Dict, List, Any, Optional
from .models import ExtractedInvoice
from .logger import logger

# Từ điển ánh xạ tự động từ tiêu đề cột trong mẫu Excel của người dùng sang trường dữ liệu hóa đơn
COLUMN_SYNONYMS: Dict[str, List[str]] = {
    "stt": ["stt", "no", "no.", "số thứ tự", "so thu tu"],
    "invoice_symbol": ["ký hiệu", "ky hieu", "mẫu số", "ký hiệu hđ", "serial", "symbol"],
    "invoice_number": ["số hóa đơn", "so hoa don", "số hđ", "invoice no", "invoice number", "số"],
    "invoice_date": ["ngày hóa đơn", "ngay hoa don", "ngày lập", "ngày tháng", "ngay", "date", "invoice date"],
    "seller_name": ["tên người bán", "ten nguoi ban", "đơn vị bán", "nhà cung cấp", "người bán", "seller", "vendor"],
    "seller_tax_id": ["mã số thuế người bán", "mst người bán", "mst bán", "mã số thuế", "mst", "tax code", "seller tax id"],
    "seller_address": ["địa chỉ người bán", "dia chi nguoi ban", "địa chỉ", "address"],
    "buyer_name": ["tên người mua", "ten nguoi mua", "đơn vị mua", "khách hàng", "người mua", "buyer"],
    "buyer_tax_id": ["mst người mua", "mã số thuế người mua", "buyer tax id"],
    "items_summary": ["nội dung", "mặt hàng", "tên hàng hóa", "diễn giải", "hàng hóa dịch vụ", "description"],
    "subtotal_amount": [
        "tiền trước thuế", "doanh số mua chưa có thuế", "giá trị chưa thuế",
        "thành tiền chưa thuế", "cộng tiền hàng", "subtotal", "amount before tax"
    ],
    "vat_rate": ["thuế suất", "thue suat", "% thuế", "vat rate", "tax rate"],
    "vat_amount": ["tiền thuế", "thuế gtgt", "tiền thuế gtgt", "vat amount", "tax amount"],
    "discount_amount": ["chiết khấu", "giảm trừ", "discount"],
    "total_amount": ["tổng tiền", "tổng thanh toán", "tổng cộng", "thành tiền", "total amount", "grand total"],
    "currency": ["loại tiền", "tiền tệ", "đvt tiền", "currency"],
    "notes": ["ghi chú", "note", "notes", "nhận xét"],
    "source_filename": ["tên file", "file nguồn", "tệp đính kèm", "source file"],
    "verification_status": ["trạng thái", "xác minh", "status"],
}


def match_header_to_field(header_text: str) -> Optional[str]:
    if not header_text:
        return None
    norm = " ".join(header_text.lower().strip().split())
    # Ưu tiên khớp chính xác trước
    for field_key, synonyms in COLUMN_SYNONYMS.items():
        for syn in synonyms:
            if norm == syn:
                return field_key
    # Sau đó kiểm tra chứa cụm từ (ưu tiên cụm từ dài hơn để tránh nhầm 'mst người mua' với 'mst')
    best_field: Optional[str] = None
    best_len = 0
    for field_key, synonyms in COLUMN_SYNONYMS.items():
        for syn in synonyms:
            if syn in norm and len(syn) > best_len:
                best_field = field_key
                best_len = len(syn)
    return best_field


def analyze_excel_template(template_path: Path) -> Dict[str, Any]:
    """
    Đọc và phân tích cấu trúc file Excel mẫu (.xlsx) do người dùng cung cấp:
    - Danh sách các sheet.
    - Hàng tiêu đề (header row) được phát hiện tự động.
    - Danh sách các cột, vị trí cột, ánh xạ trường hóa đơn gợi ý.
    - Phát hiện các cột có sẵn công thức ở dòng mẫu đầu tiên để bảo toàn không ghi đè.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Chưa cài đặt thư viện openpyxl. Hãy chạy: pip install openpyxl==3.1.5"
        ) from exc

    if not template_path.exists():
        raise FileNotFoundError("Không tìm thấy file Excel mẫu.")

    wb = openpyxl.load_workbook(filename=str(template_path), data_only=False)
    sheets_info = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        best_header_row = 1
        best_match_count = -1
        detected_columns = []

        # Quét 20 dòng đầu tiên để tìm dòng tiêu đề bảng kê
        max_scan_row = min(ws.max_row or 1, 20)
        max_scan_col = min(ws.max_column or 1, 40)

        for r in range(1, max_scan_row + 1):
            row_cols = []
            distinct_fields = set()
            distinct_texts = set()
            for c in range(1, max_scan_col + 1):
                val = ws.cell(row=r, column=c).value
                text = str(val).strip() if val is not None else ""
                if text and not text.startswith("="):
                    distinct_texts.add(text)
                    mapped = match_header_to_field(text)
                    if mapped:
                        distinct_fields.add(mapped)
                    row_cols.append(
                        {
                            "col_index": c,
                            "col_letter": openpyxl.utils.get_column_letter(c),
                            "header_text": text,
                            "mapped_field": mapped,
                        }
                    )
            if len(distinct_fields) > best_match_count and len(distinct_texts) >= 2:
                best_match_count = len(distinct_fields)
                best_header_row = r
                detected_columns = row_cols

        # Kiểm tra dòng ngay dưới dòng tiêu đề xem có chứa công thức mẫu không
        sample_data_row = best_header_row + 1
        formula_columns = []
        for col_info in detected_columns:
            c_idx = col_info["col_index"]
            sample_val = ws.cell(row=sample_data_row, column=c_idx).value
            if isinstance(sample_val, str) and sample_val.startswith("="):
                col_info["has_formula"] = True
                col_info["formula_template"] = sample_val
                formula_columns.append(col_info["col_letter"])
            else:
                col_info["has_formula"] = False
                col_info["formula_template"] = None

        sheets_info.append(
            {
                "sheet_name": sheet_name,
                "header_row": best_header_row,
                "start_data_row": best_header_row + 1,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "columns": detected_columns,
                "formula_columns": formula_columns,
            }
        )

    wb.close()
    return {
        "template_filename": template_path.name,
        "sheet_names": [s["sheet_name"] for s in sheets_info],
        "active_sheet": sheets_info[0]["sheet_name"] if sheets_info else "",
        "sheets": sheets_info,
    }


def _copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format


def export_invoices_to_template(
    template_path: Path,
    output_path: Path,
    invoices: List[ExtractedInvoice],
    target_sheet: Optional[str] = None,
    custom_column_mapping: Optional[Dict[int, str]] = None,
    include_audit_sheet: bool = True,
) -> Path:
    """
    Ghi danh sách hóa đơn vào đúng sheet và đúng cột của file Excel mẫu.
    - Giữ nguyên cấu trúc các sheet, dòng tiêu đề, định dạng ô (font, viền, number_format).
    - Không ghi đè làm hỏng các cột chứa công thức Excel.
    - Tùy chọn thêm sheet 'Bao_Cao_Kiem_Tra' ghi lại cảnh báo và lịch sử chỉnh sửa.
    """
    if not invoices:
        raise ValueError("Không có dữ liệu hóa đơn hợp lệ để xuất Excel.")
    if not template_path.exists():
        raise FileNotFoundError("Người dùng chưa cung cấp file Excel mẫu.")

    import openpyxl

    analysis = analyze_excel_template(template_path)
    wb = openpyxl.load_workbook(filename=str(template_path))

    sheet_meta = None
    if target_sheet and target_sheet in wb.sheetnames:
        sheet_meta = next((s for s in analysis["sheets"] if s["sheet_name"] == target_sheet), None)
    if sheet_meta is None:
        sheet_meta = analysis["sheets"][0]

    ws = wb[sheet_meta["sheet_name"]]
    header_row = sheet_meta["header_row"]
    start_row = sheet_meta["start_data_row"]

    # Xây dựng bảng ánh xạ cột -> trường dữ liệu
    col_map: Dict[int, str] = {}
    for col in sheet_meta["columns"]:
        c_idx = col["col_index"]
        if custom_column_mapping and c_idx in custom_column_mapping:
            col_map[c_idx] = custom_column_mapping[c_idx]
        elif col["mapped_field"]:
            col_map[c_idx] = col["mapped_field"]

    # Lưu lại style và công thức mẫu của dòng dữ liệu đầu tiên
    style_reference_row = start_row
    for idx, inv in enumerate(invoices):
        current_row = start_row + idx
        for col in sheet_meta["columns"]:
            c_idx = col["col_index"]
            ref_cell = ws.cell(row=style_reference_row, column=c_idx)
            target_cell = ws.cell(row=current_row, column=c_idx)

            if current_row != style_reference_row:
                _copy_cell_style(ref_cell, target_cell)

            # Nếu cột có công thức ở dòng mẫu -> dịch chuyển công thức theo số dòng
            if col.get("has_formula") and col.get("formula_template"):
                formula_str = str(col["formula_template"])
                translated_formula = formula_str.replace(str(style_reference_row), str(current_row))
                target_cell.value = translated_formula
                continue

            field_name = col_map.get(c_idx)
            if not field_name:
                continue

            if field_name == "stt":
                target_cell.value = idx + 1
            else:
                val = getattr(inv, field_name, "")
                target_cell.value = val if val is not None else ""

    # Tùy chọn tạo thêm sheet Báo cáo kiểm tra & truy nguyên
    if include_audit_sheet:
        audit_name = "Bao_Cao_Kiem_Tra"
        if audit_name in wb.sheetnames:
            del wb[audit_name]
        audit_ws = wb.create_sheet(title=audit_name)
        audit_headers = [
            "STT",
            "Tên file PDF",
            "Số hóa đơn",
            "MST Người bán",
            "Phương thức đọc",
            "Trạng thái xác minh",
            "Trường cần lưu ý",
            "Chi tiết cảnh báo / Kiểm tra",
        ]
        audit_ws.append(audit_headers)
        for idx, inv in enumerate(invoices):
            issues_str = "; ".join(f"[{i.severity}] {i.message}" for i in inv.validation_issues)
            uncertain_str = ", ".join(inv.uncertain_fields)
            audit_ws.append(
                [
                    idx + 1,
                    inv.source_filename,
                    inv.invoice_number,
                    inv.seller_tax_id,
                    inv.extraction_method,
                    inv.verification_status,
                    uncertain_str or "Không",
                    issues_str or "Đạt kiểm tra",
                ]
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    wb.close()
    logger.info("Đã xuất thành công %d hóa đơn ra file Excel: %s", len(invoices), output_path.name)
    return output_path
